"""Exportación a Excel con formato."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from nucleo import PRIORIDADES

MONEDA = ("Precio unitario", "Monto", "Monto mínimo", "Total", "Presupuesto", "Presupuesto utilizado",
          "Saldo", "Escenario mínimo", "Saldo con mínimo", *PRIORIDADES)


def _libro(hojas: dict[str, pd.DataFrame]) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for hoja, data in hojas.items():
            data.to_excel(xw, sheet_name=hoja, index=False)
            ws = xw.sheets[hoja]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F7A4D")
                c.alignment = Alignment(vertical="center", wrap_text=True)
            for j, col in enumerate(data.columns, start=1):
                largo = max([len(str(col))] + [len(str(v)) for v in data[col].head(200)])
                ws.column_dimensions[get_column_letter(j)].width = min(max(10, largo + 2), 55)
                if col in MONEDA:
                    for celdas in ws.iter_cols(min_col=j, max_col=j, min_row=2):
                        for c in celdas:
                            c.number_format = '"$" #,##0.00'
    return buf.getvalue()


def registros(df: pd.DataFrame) -> bytes:
    """Detalle de registros + resumen por programa y prioridad."""
    nombres = {
        "prog": "Programa", "partida_codigo": "Cód. partida", "partida": "Partida",
        "codigo": "Código", "item": "Ítem",
        "objeto_gasto": "Objeto gasto", "clasificador": "Clasificador", "cantidad": "Cantidad",
        "cantidad_minima": "Cant. mínima", "unidad": "Unidad", "precio_unitario": "Precio unitario",
        "monto": "Monto", "monto_minimo": "Monto mínimo", "prioridad": "Prioridad",
        "justificacion": "Justificación", "creado_en": "Fecha de carga", "cargado_por": "Cargado por",
    }
    detalle = df.drop(columns=["id", "nro_programa", "programa", "jurisdiccion"]).rename(columns=nombres)
    detalle = detalle[["Programa"] + [c for c in detalle.columns if c != "Programa"]]
    resumen = (df.pivot_table(index="prog", columns="prioridad", values="monto", aggfunc="sum", fill_value=0)
                 .reindex(columns=PRIORIDADES, fill_value=0))
    resumen["Total"] = resumen.sum(axis=1)
    resumen["Monto mínimo"] = df.groupby("prog")["monto_minimo"].sum()
    resumen = resumen.reset_index().rename(columns={"prog": "Programa"})
    return _libro({"Detalle": detalle, "Resumen": resumen})


def tabla(df: pd.DataFrame, hoja: str = "Datos") -> bytes:
    return _libro({hoja: df})
